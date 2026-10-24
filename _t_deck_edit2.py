"""Workplan deck edit (blocked until PowerPoint closes the file):
adds a GraphRAG-backend band to slide 7 (FAQ path) of
Kallang_Track2_Workplan_Architecture_Workflow.pptx, keeping the TF-IDF
pipeline shown as the judged default. Run from track2_service_chatbot/.
"""
import copy

from pptx import Presentation
from pptx.util import Inches, Pt

PATH = "docs/Kallang_Track2_Workplan_Architecture_Workflow.pptx"

BAND = (
    "Optional backend (sidebar, live 24 Oct): GraphRAG — a code router picks local / global / hybrid; "
    "sweep questions (“summarise all your policies”) answer from community breadth with a multi-section "
    "Source. TF-IDF + synonym map stays the default and the abstention backstop. "
    "G8 re-run: hit@1 78% · abstention 100% · global coverage 6.3 vs 3.7 sections."
)

try:
    open(PATH, "r+b").close()
except PermissionError:
    raise SystemExit("STILL OPEN IN POWERPOINT — close it and re-run this script.")

prs = Presentation(PATH)
slide = prs.slides[6]

# find the helper line "score < τ → …" to steal its font attributes
src_run = None
for sh in slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith("score < τ"):
        for para in sh.text_frame.paragraphs:
            if para.runs:
                src_run = para.runs[0]
                break
    if src_run:
        break

box = slide.shapes.add_textbox(Inches(0.62), Inches(6.45), Inches(12.09), Inches(0.55))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run()
r.text = BAND
if src_run is not None:
    r.font.size = src_run.font.size or Pt(10)
    r.font.name = src_run.font.name
    r.font.italic = True
    if src_run.font.color and src_run.font.color.type is not None:
        r.font.color.rgb = src_run.font.color.rgb
else:
    r.font.size = Pt(10)
    r.font.italic = True

prs.save(PATH)
print("band added to slide 7; saved", PATH)
